# -*- coding: utf-8 -*-
"""生成 issue-529 复现数据: 含 TOC 目录域的 docx(条目文本为普通 run)

上游 issue #529 报告打开 word 文件时目录仅显示序号(页码), 条目内容空白。
构造与 issue #597 plain 变体同类的结构: TOC 域结果区条目为普通文本 run +
嵌套 PAGEREF 域(无 <w:hyperlink> 包裹), TOC 域 begin/end 分处首末条目段。
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "word", "test-word-cweijan-529-toc-blank.docx")

TAB_POS = "9350"

TOC_STYLES = (
    '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:style w:type="paragraph" w:styleId="TOCHeading"><w:name w:val="TOC Heading"/>'
    '<w:basedOn w:val="Normal"/><w:rPr><w:b/><w:sz w:val="32"/></w:rPr></w:style>'
    '<w:style w:type="paragraph" w:styleId="TOC1"><w:name w:val="toc 1"/><w:basedOn w:val="Normal"/>'
    f'<w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{TAB_POS}"/></w:tabs></w:pPr></w:style>'
    '<w:style w:type="paragraph" w:styleId="TOC2"><w:name w:val="toc 2"/><w:basedOn w:val="Normal"/>'
    f'<w:pPr><w:ind w:left="220"/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{TAB_POS}"/></w:tabs></w:pPr></w:style>'
    '</w:styles>'
)


def r_text(t):
    return f'<w:r><w:t xml:space="preserve">{t}</w:t></w:r>'


def r_tab():
    return '<w:r><w:tab/></w:r>'


def r_field(char):
    return f'<w:r><w:fldChar w:fldCharType="{char}"/></w:r>'


def r_instr(instr):
    return f'<w:r><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'


def pageref(anchor, page):
    """嵌套 PAGEREF 域: begin instr separate 页码 end"""
    return (r_field("begin") + r_instr(f'PAGEREF {anchor} \\h') + r_field("separate")
            + r_text(page) + r_field("end"))


def toc_paragraph(text, anchor, page, style, toc_begin=False, toc_end=False):
    parts = []
    if toc_begin:
        parts += [r_field("begin"), r_instr('TOC \\o "1-2" \\z \\u'), r_field("separate")]
    parts += [r_text(text), r_tab(), pageref(anchor, page)]
    if toc_end:
        parts += [r_tab(), r_field("end")]
    return f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr>{"".join(parts)}</w:p>'


def main():
    doc = Document()

    entries = [
        ("_Toc52901", "第一章 项目概述", "1", "TOC1"),
        ("_Toc52902", "第二章 需求分析", "2", "TOC1"),
        ("_Toc52903", "2.1 功能需求", "2", "TOC2"),
        ("_Toc52904", "2.2 非功能需求", "3", "TOC2"),
        ("_Toc52905", "第三章 系统设计", "4", "TOC1"),
        ("_Toc52906", "第四章 总结与展望", "5", "TOC1"),
    ]
    paras = ['<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>目录</w:t></w:r></w:p>']
    for i, (anchor, text, page, style) in enumerate(entries):
        paras.append(toc_paragraph(text, anchor, page, style,
                                   toc_begin=(i == 0), toc_end=(i == len(entries) - 1)))
    toc_block = parse_xml(
        '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/></w:docPartObj></w:sdtPr>'
        f'<w:sdtContent>{"".join(paras)}</w:sdtContent></w:sdt>')
    doc.element.body.insert(0, toc_block)
    for st in parse_xml(TOC_STYLES):
        doc.styles.element.append(st)

    def heading(text, level, name, bid):
        p = doc.add_heading(text, level=level)
        p._p.insert(1, parse_xml(
            f'<w:bookmarkStart xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:id="{bid}" w:name="{name}"/>'))
        p._p.append(parse_xml(
            f'<w:bookmarkEnd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:id="{bid}"/>'))

    heading("第一章 项目概述", 1, "_Toc52901", 300)
    doc.add_paragraph("第一章正文(issue-529: 目录条目文本为普通 run 的 TOC 域)。")
    doc.paragraphs[-1].runs[0].add_break(WD_BREAK.PAGE)
    heading("第二章 需求分析", 1, "_Toc52902", 301)
    doc.add_paragraph("第二章正文。")
    heading("2.1 功能需求", 2, "_Toc52903", 302)
    doc.add_paragraph("2.1 正文。")
    doc.paragraphs[-1].runs[0].add_break(WD_BREAK.PAGE)
    heading("2.2 非功能需求", 2, "_Toc52904", 303)
    doc.add_paragraph("2.2 正文。")
    doc.paragraphs[-1].runs[0].add_break(WD_BREAK.PAGE)
    heading("第三章 系统设计", 1, "_Toc52905", 304)
    doc.add_paragraph("第三章正文。")
    doc.paragraphs[-1].runs[0].add_break(WD_BREAK.PAGE)
    heading("第四章 总结与展望", 1, "_Toc52906", 305)
    doc.add_paragraph("第四章正文。")

    doc.save(OUT)
    print("生成:", OUT)


if __name__ == "__main__":
    main()
