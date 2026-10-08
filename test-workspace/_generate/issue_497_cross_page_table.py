# -*- coding: utf-8 -*-
"""生成 issue-497 复现数据: 表格跨页的 docx

上游 issue #497 报告"跨页面的表格会显示不完整"(正文仅一张截图)。
按"跨页面"的分页特征归为 Word 预览: 生成 100 行 x 4 列长表格,
首行设置跨页重复表头, 每第 15 行设置较大行高使行边界与分页边界交错,
表格前后各有正文段落便于对比。
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "word", "test-word-cweijan-497-cross-page-table.docx")

ROWS = 100


def set_repeat_header(row):
    """设置首行跨页重复表头(w:tblHeader)"""
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)


def main():
    doc = Document()
    doc.add_heading("跨页表格显示测试 (issue #497)", 0)
    doc.add_paragraph(f"下方长表格共 {ROWS} 行, 在 Word/WPS 中跨约 3 页, 用于验证跨页表格是否完整渲染。")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, t in enumerate(["序号", "名称", "说明", "备注"]):
        table.rows[0].cells[i].text = t
    set_repeat_header(table.rows[0])

    for r in range(1, ROWS + 1):
        cells = table.add_row().cells
        cells[0].text = str(r)
        cells[1].text = f"条目 {r:03d}"
        cells[2].text = f"第 {r} 行说明文本, 用于占位并测试跨页显示完整性。"
        cells[3].text = "跨页标记" if r % 7 == 0 else ""
        if r % 15 == 0:  # 少量高行, 让行边界与分页边界交错
            row = table.rows[-1]
            row.height = Cm(2.5)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    doc.add_paragraph("表格之后的正文段落, 用于确认跨页表格之后的内容渲染正常。")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("生成:", OUT)


if __name__ == "__main__":
    main()
