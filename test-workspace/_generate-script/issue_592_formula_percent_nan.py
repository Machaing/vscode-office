# -*- coding: utf-8 -*-
"""生成 issue-592 复现数据: 含 $ 绝对引用公式 + 缓存值的 xlsx

上游 issue #592: E6=23/E7=40/E8=21, E9==SUM(E6:E8), F6==E6/$E$9(格式 0.0%),
文件内缓存 F6=0.27380952380952384; 查看器显示 NaN%。
openpyxl 保存公式时不写缓存值, 故先生成后用 zipfile 后处理,
向公式单元格注入与公式一致的 <v> 缓存结果, 模拟 Excel/WPS 保存的真实文件。
"""
import os
import re
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8")

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "excel", "test-excel-cweijan-592-formula-percent-nan.xlsx")
TMP = os.path.join(ROOT, "_generate-script", "_592_tmp.xlsx")


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "占比"

    ws["C5"] = "单价"
    ws["D5"] = "标签"
    ws["E5"] = "数值"
    ws["F5"] = "占比(=E/$E$9)"
    ws["G5"] = "全绝对引用"

    ws["C6"] = 2.5
    ws["E6"], ws["E7"], ws["E8"] = 23, 40, 21
    ws["E9"] = "=SUM(E6:E8)"
    for r in (6, 7, 8):
        ws[f"F{r}"] = f"=E{r}/$E$9"
        ws[f"F{r}"].number_format = "0.0%"
    ws["F9"] = "=SUM(F6:F8)"  # 间接引用含 $ 公式的单元格
    ws["F9"].number_format = "0.0%"
    ws["D15"] = "数量"
    ws["E15"] = 10
    ws["G15"] = "=$E$15*$C$6"  # issue 提到的全绝对引用形态

    for col, width in zip("CDEFG", [10, 10, 10, 16, 12]):
        ws.column_dimensions[col].width = width

    # 缓存结果与公式保持一致
    total = 23 + 40 + 21  # 84
    cached = {
        "E9": total,
        "F6": 23 / total,
        "F7": 40 / total,
        "F8": 21 / total,
        "F9": 1.0,
        "G15": 25,
    }

    wb.save(TMP)

    def inject(xml: str) -> str:
        for ref, value in cached.items():
            # openpyxl 对公式单元格写 <f>...</f><v></v>, 将空 <v> 替换为缓存结果(无 <v> 时补写)
            pat = re.compile(r'(<c r="%s"[^>]*>)(<f>[^<]*</f>)(?:<v>[^<]*</v>)?(</c>)' % ref)
            xml, n = pat.subn(lambda m: f"{m.group(1)}{m.group(2)}<v>{value!r}</v>{m.group(3)}", xml)
            if n != 1:
                raise RuntimeError(f"公式单元格 {ref} 注入缓存值失败(匹配 {n} 次)")
        return xml

    with zipfile.ZipFile(TMP) as zin:
        items = [(info, zin.read(info.filename)) for info in zin.infolist()]
    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename == "xl/worksheets/sheet1.xml":
                data = inject(data.decode("utf-8")).encode("utf-8")
            zout.writestr(info, data)
    os.remove(TMP)
    print("生成:", OUT)
    print("缓存值:", {k: repr(v) for k, v in cached.items()})


if __name__ == "__main__":
    main()
