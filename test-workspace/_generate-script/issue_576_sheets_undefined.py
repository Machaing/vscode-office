# -*- coding: utf-8 -*-
"""生成 issue-576 复现数据: xl/workbook.xml 异常的 xlsx

上游 issue #576 报告打开部分 xlsx 直接报
"Cannot read properties of undefined (reading 'sheets')"。
扩展对 xlsx 走 ExcelJS 加载, @cweijan/exceljs 的 xlsx/xlsx.js 解析到
xl/workbook.xml 时执行 `model.sheets = workbook.sheets`; 当 WorkbookXform
无法从该 XML 解析出 model(如文件为空/根元素不是 <workbook>)时,
parseWorkbook 返回 undefined, 读取 workbook.sheets 即抛出该错误。
这里先用 openpyxl 生成正常 xlsx, 再把包内 xl/workbook.xml 改写为空文件
重新打包, 其余部件均保持完好, 作为确定触发同类报错的最小构造。
"""
import os
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8")

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "excel", "test-excel-cweijan-576-sheets-undefined.xlsx")
TMP = os.path.join(ROOT, "_generate-script", "_576_tmp.xlsx")


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "issue-576"
    ws["B1"] = "workbook.xml 为空的异常 xlsx"
    for r in range(2, 6):
        ws.cell(row=r, column=1, value=f"row-{r}")
        ws.cell(row=r, column=2, value=r * 10)
    wb.save(TMP)

    # 重新打包: 仅将 xl/workbook.xml 置空, 其余部件原样保留
    with zipfile.ZipFile(TMP) as zin:
        items = [(info, zin.read(info.filename)) for info in zin.infolist()]
    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename == "xl/workbook.xml":
                data = b""
            zout.writestr(info, data)
    os.remove(TMP)
    print("生成:", OUT)


if __name__ == "__main__":
    main()
