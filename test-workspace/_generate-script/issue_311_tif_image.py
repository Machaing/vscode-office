# -*- coding: utf-8 -*-
"""生成 issue-311 复现数据: docx 内嵌 TIFF 图片

上游 issue #311 报告 Word 文档中的 .tif 图片无法预览。
用 Pillow 生成同一张测试图, 分别以 PNG(对照)与 TIFF(复现)内嵌到 docx。
"""
import io
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Cm
from PIL import Image, ImageDraw

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "word", "test-word-cweijan-311-tif-image.docx")

TOP = (66, 133, 244)
BOTTOM = (15, 157, 88)


def make_tif_bytes(size=320):
    """生成带渐变背景 + 同心圆 + 对角线的 TIFF(RGB 基线, 未压缩)"""
    img = Image.new("RGB", (size, size), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    for y in range(size):
        t = y / size
        draw.line([(0, y), (size, y)],
                  fill=tuple(int(a + (b - a) * t) for a, b in zip(TOP, BOTTOM)))
    for i in range(6):
        r = size // 2 - i * 24
        if r > 0:
            draw.ellipse([size // 2 - r, size // 2 - r, size // 2 + r, size // 2 + r],
                         outline=(219, 68, 55), width=5)
    draw.line([(0, 0), (size, size)], fill=(244, 180, 0), width=4)
    draw.line([(size, 0), (0, size)], fill=(255, 255, 255), width=3)
    buf = io.BytesIO()
    img.save(buf, format="TIFF")
    return buf.getvalue()


def main():
    tif_bytes = make_tif_bytes()
    # 同一张图转 PNG 作为对照组
    png_buf = io.BytesIO()
    Image.open(io.BytesIO(tif_bytes)).save(png_buf, format="PNG")

    doc = Document()
    doc.add_heading("TIFF 图片预览测试 (issue #311)", 0)
    doc.add_paragraph("下面两张图为同一内容, 分别以 PNG 与 TIFF 内嵌, 用于对照预览效果。")

    doc.add_heading("1. PNG 图片(对照组)", level=1)
    doc.add_picture(io.BytesIO(png_buf.getvalue()), width=Cm(8))
    doc.add_paragraph("预期: PNG 正常显示。")

    doc.add_heading("2. TIFF 图片(复现组)", level=1)
    doc.add_picture(io.BytesIO(tif_bytes), width=Cm(8))
    doc.add_paragraph("issue 现象: TIFF 图片无法预览(空白/不渲染)。")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("生成:", OUT)


if __name__ == "__main__":
    main()
