# -*- coding: utf-8 -*-
"""验证生成的测试文档可被相应解析器读取"""
import os
import sys
import tarfile
import zipfile

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ok, fail = [], []


def check(name, fn):
    try:
        fn()
        ok.append(name)
        print(f"  PASS {name}")
    except Exception as e:
        fail.append((name, str(e)))
        print(f"  FAIL {name}: {e}")


# 图片类: PIL 打开并解码
def verify_images():
    from PIL import Image
    for f in sorted(os.listdir(os.path.join(ROOT, "image"))):
        with Image.open(os.path.join(ROOT, "image", f)) as im:
            im.load()


def verify_tiff():
    from PIL import Image
    with Image.open(os.path.join(ROOT, "tiff", "sample.tiff")) as im:
        assert im.n_frames == 3, f"期望 3 页, 实际 {im.n_frames}"


def verify_icns():
    data = open(os.path.join(ROOT, "icns", "sample.icns"), "rb").read()
    assert data[:4] == b"icns", "magic 错误"
    size = int.from_bytes(data[4:8], "big")
    assert size == len(data), "总长度字段错误"


# zip 容器类: zipfile 枚举
def zip_names(rel):
    def fn():
        with zipfile.ZipFile(os.path.join(ROOT, rel)) as z:
            names = z.namelist()
            assert len(names) > 0, "空压缩包"
    return fn


def verify_epub():
    with zipfile.ZipFile(os.path.join(ROOT, "epub", "sample.epub")) as z:
        infos = z.infolist()
        assert infos[0].filename == "mimetype" and infos[0].compress_type == zipfile.ZIP_STORED, "mimetype 必须第一个且不压缩"
        assert z.read("mimetype") == b"application/epub+zip"
        assert "META-INF/container.xml" in z.namelist()


def verify_tar():
    with tarfile.open(os.path.join(ROOT, "archive", "sample.tar")) as t:
        assert len(t.getnames()) >= 2


def verify_tgz():
    with tarfile.open(os.path.join(ROOT, "archive", "sample.tar.gz"), "r:gz") as t:
        assert len(t.getnames()) >= 2


def verify_crx():
    data = open(os.path.join(ROOT, "archive", "sample.crx"), "rb").read()
    assert data[:4] == b"Cr24" and int.from_bytes(data[4:8], "little") == 2, "CRX2 头错误"
    with zipfile.ZipFile(io.BytesIO(data[16 + int.from_bytes(data[8:12], "little") + int.from_bytes(data[12:16], "little"):])) as z:
        assert len(z.namelist()) > 0


import io


def verify_xlsx():
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(ROOT, "excel", "sample.xlsx"))
    assert wb.sheetnames == ["销售数据", "说明"]
    ws = wb["销售数据"]
    assert ws["E2"].value == "=SUM(B2:D2)"


def verify_docx():
    from docx import Document
    doc = Document(os.path.join(ROOT, "word", "sample.docx"))
    assert doc.paragraphs[0].style.name == "Title"
    assert len(doc.tables) == 1


def verify_pptx():
    from pptx import Presentation
    prs = Presentation(os.path.join(ROOT, "powerpoint", "sample.pptx"))
    assert len(prs.slides.__iter__.__self__._sldIdLst) >= 3 or len(prs.slides._sldIdLst) >= 3


def verify_class():
    for f in os.listdir(os.path.join(ROOT, "java")):
        data = open(os.path.join(ROOT, "java", f), "rb").read()
        assert data[:4] == b"\xCA\xFE\xBA\xBE", f"{f} 不是有效 class 文件"


def verify_fonts():
    for f in ["sample.ttf", "sample.woff", "sample.woff2", "sample.otf"]:
        p = os.path.join(ROOT, "font", f)
        data = open(p, "rb").read()
        assert len(data) > 1000, f"{f} 过小"
        if f.endswith("ttf") or f.endswith("otf"):
            assert data[:4] in (b"\x00\x01\x00\x00", b"OTTO"), f"{f} sfnt 头错误"
        elif f.endswith("woff"):
            assert data[:4] == b"wOFF", f"{f} woff 头错误"
        else:
            assert data[:4] == b"wOF2", f"{f} woff2 头错误"


def verify_psd():
    from PIL import Image
    with Image.open(os.path.join(ROOT, "psd", "sample.psd")) as im:
        assert im.size == (128, 128), f"尺寸错误 {im.size}"
        im.load()


def verify_svg():
    data = open(os.path.join(ROOT, "svg", "sample.svg"), encoding="utf-8").read()
    assert "<svg" in data and "</svg>" in data


if __name__ == "__main__":
    print("== Python 侧验证 ==")
    check("image/* (PIL 全部解码)", verify_images)
    check("tiff 多页", verify_tiff)
    check("psd (PIL)", verify_psd)
    check("icns 结构", verify_icns)
    check("epub 结构", verify_epub)
    check("zip/jar/apk/vsix", zip_names("archive/sample.zip"))
    check("jar", zip_names("archive/sample.jar"))
    check("apk", zip_names("archive/sample.apk"))
    check("vsix", zip_names("archive/sample.vsix"))
    check("crx 头", verify_crx)
    check("tar", verify_tar)
    check("tar.gz/tgz", verify_tgz)
    check("xlsx (openpyxl)", verify_xlsx)
    check("docx (python-docx)", verify_docx)
    check("pptx (python-pptx)", verify_pptx)
    check("class 文件头", verify_class)
    check("font 魔数", verify_fonts)
    check("svg 文本", verify_svg)
    print(f"\n通过 {len(ok)}, 失败 {len(fail)}")
    sys.exit(1 if fail else 0)
