# -*- coding: utf-8 -*-
"""生成 issue-607 复现数据: 无 <dimension> 元素 + x: 命名空间前缀的 xlsx

上游 issue #607: OOXML 导出器生成的 xlsx 打开后仅显示空白 Sheet1,
无 sheet 名、无数据、无报错; 常规文件正常。

根因二分结论(见 docs/plans/todo/plan-cweijan-issues-607-missing-dimension.md):
- 仅删除 <dimension> 不触发 bug: ExcelJS 用 eachRow 遍历实际行, SheetJS 缺
  dimension 时自动从 sheetData 推断 !ref, 两边都不依赖该元素;
- 真正根因是 worksheet XML 使用 x: 命名空间前缀: ExcelJS xform 按无前缀
  标签名(worksheet/sheetData/row/c)匹配, x: 前缀导致 parseStream 返回
  undefined, 随后 worksheet.sheetNo 赋值抛
  "TypeError: Cannot set properties of undefined (setting 'sheetNo')";
- content type 差异(Default 用 workbook 主类型、删 /xl/workbook.xml override)
  无影响: ExcelJS 按 zip 内路径硬编码定位 part, 不读 [Content_Types].xml。

本仓库已由 issue-576 修复(ExcelJS 解析抛错时降级 SheetJS 兜底)顺带覆盖:
SheetJS 对 x: 前缀解析正常。两个复现文件分别用于回归验证:
1. missing-dimension: 删 <dimension>(排除项, 打开应正常);
2. x-prefixed-ns: 删 <dimension> + worksheet 元素全部改 x: 前缀(根因项,
   走 ExcelJS 抛错 → SheetJS 兜底链路, 打开应正常显示数据)。
"""
import os
import re
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8")

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "excel")
TMP = os.path.join(ROOT, "_generate-script", "_607_tmp.xlsx")

DIM_PAT = re.compile(r"<dimension[^>]*/>")
TAG_PAT = re.compile(r"<(\/?)([a-zA-Z][\w.-]*)")


def build_workbook() -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Data1"
    ws1.append(["日期", "数值", "文本"])
    for i in range(1, 21):
        ws1.append([f"2026-09-{i:02d}", i * 1.5, f"row-{i}"])

    ws2 = wb.create_sheet("Data2")
    ws2.append(["项目", "数量"])
    for i in range(1, 11):
        ws2.append([f"item-{i}", i])

    for ws in (ws1, ws2):
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)

    wb.save(TMP)


def strip_dimension(xml: str) -> str:
    xml, n = DIM_PAT.subn("", xml)
    if n != 1:
        raise RuntimeError(f"dimension 元素匹配 {n} 次(预期 1)")
    return xml


def add_x_prefix(xml: str) -> str:
    """所有元素名加 x: 前缀, 根元素 xmlns= 改为 xmlns:x=(复刻 issue 中导出器行为)"""
    xml = TAG_PAT.sub(lambda m: f"<{m.group(1)}x:{m.group(2)}", xml)
    return re.sub(r"(\s)xmlns=", r"\1xmlns:x=", xml)


def rewrite(out_path: str, transform) -> None:
    with zipfile.ZipFile(TMP) as zin:
        items = [(info, zin.read(info.filename)) for info in zin.infolist()]
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", info.filename):
                data = transform(data.decode("utf-8")).encode("utf-8")
            zout.writestr(info, data)


def main():
    build_workbook()

    out1 = os.path.join(OUT_DIR, "test-excel-cweijan-607-missing-dimension.xlsx")
    rewrite(out1, strip_dimension)
    print("生成:", out1)
    print("  已删除全部 worksheet 的 <dimension> 元素(排除项, 打开应正常)")

    out2 = os.path.join(OUT_DIR, "test-excel-cweijan-607-x-prefixed-ns.xlsx")
    rewrite(out2, lambda xml: add_x_prefix(strip_dimension(xml)))
    print("生成:", out2)
    print("  删 <dimension> + worksheet 元素全部改 x: 前缀(根因项, 验证 SheetJS 兜底)")

    os.remove(TMP)


if __name__ == "__main__":
    main()
